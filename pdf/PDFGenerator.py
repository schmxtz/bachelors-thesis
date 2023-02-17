from utils.Files import parse_excel_file, PLACEHOLDER_CLOSING, PLACEHOLDER_OPENING
from utils.Util import has_numbers
import os, time, subprocess, logging
from threading import Thread
from docx import Document

DOCX_EXT = '.docx'
EXCEL_EXT = '.xlsx'


class PDFGenerator:
    def __init__(self, template_file_name: str, excel_file_name: str, output_path: str, delete_source_docx: bool = True):
        if not template_file_name.endswith(DOCX_EXT):
            raise ValueError('Template file must be {} format.'.format(DOCX_EXT))
        self.template_file_name = template_file_name

        if not excel_file_name.endswith(EXCEL_EXT):
            raise ValueError('Excel sheet file must be {} format.'.format(EXCEL_EXT))
        self.excel_file_name = excel_file_name

        if not os.path.isdir(output_path):
            raise ValueError('Output path must be a directory.')
        self.output_path = output_path

        self.delete_source_docx = delete_source_docx
        self.parameters = None

    def parse_excel_file(self):
        self.parameters = parse_excel_file(self.excel_file_name)

    def replace_parameters(self):
        logging.info('Creating word files... (0/{0})'.format(len(self.parameters)))
        if self.parameters is None:
            return
        ctr = 1
        len_params = len(self.parameters)
        for parameter in self.parameters:
            logging.info('Creating word files... ({0}/{1})'.format(ctr, len_params))
            template = Document(self.template_file_name)
            file_name = self.build_file_name(parameter)
            for paragraph in template.paragraphs:
                self.replace_text_in_paragraph(paragraph, parameter)

            for table in template.tables:
                for col in table.columns:
                    for cell in col.cells:
                        for paragraph in cell.paragraphs:
                            self.replace_text_in_paragraph(paragraph, parameter)
            template.save(self.output_path + '/' + file_name)
            ctr += 1
        logging.info('Finished creating word files')

    def convert_docx_to_pdf(self):
        current_wcd = os.getcwd()
        current_wcd = current_wcd.replace('\\', '/')
        command = 'wscript {0}/pdf/DocxToPdf.vbs {1} {2}'.format(current_wcd,
                                                                 self.output_path,
                                                                 int(self.delete_source_docx))
        subprocess.call(command)
        logging.info('Finished converting PDFs')

    def convert_docx_to_pdf_thread(self):
        Thread(target=self.convert_docx_to_pdf).start()

    @staticmethod
    def replace_text_in_paragraph(paragraph, parameter):
        line = paragraph.runs
        for word in line:
            if PLACEHOLDER_OPENING in word.text and PLACEHOLDER_CLOSING in word.text:
                if word.text in parameter:
                    value = parameter.get(word.text)
                    if value is None:
                        value = ''
                    word.text = word.text.replace(word.text, value)
                else:
                    word.text = word.text.replace(word.text, '')

    @staticmethod
    def delete_paragraph(paragraph):
        p = paragraph._element
        p.getparent().remove(p)
        paragraph._p = paragraph._element = None

    @staticmethod
    def build_file_name(parameter):
        return '{last_name}_{first_name}_{module}_{time}.docx'.format(
            last_name=parameter[PLACEHOLDER_OPENING + 'Nachname' + PLACEHOLDER_CLOSING],
            first_name=parameter[PLACEHOLDER_OPENING + 'Vorname' + PLACEHOLDER_CLOSING],
            module=parameter[PLACEHOLDER_OPENING + 'Modul' + PLACEHOLDER_CLOSING],
            time=time.time_ns())

