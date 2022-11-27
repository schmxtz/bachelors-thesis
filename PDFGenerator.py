from docx import Document
from utils.Files import parse_excel_file, PLACEHOLDER_CLOSING, PLACEHOLDER_OPENING
import os, time
import subprocess


class PDFGenerator:
    def __init__(self, template_file_name: str, excel_file_name: str, output_path: str):
        if not template_file_name.endswith('.docx'):
            raise ValueError('Template file must be .docx format.')
        self.template_file_name = template_file_name
        self.output_path = output_path
        self.parameters, self.header = parse_excel_file(excel_file_name)

    def replace_parameters(self):
        for parameter in self.parameters:
            template = Document(self.template_file_name)
            file_name = self.build_file_name(parameter)
            for paragraph in template.paragraphs:
                self.replace_text_in_paragraph(paragraph, parameter)

            for table in template.tables:
                for col in table.columns:
                    for cell in col.cells:
                        for paragraph in cell.paragraphs:
                            self.replace_text_in_paragraph(paragraph, parameter)
            template.save(self.output_path + file_name)

    def convert_docx_to_pdf(self):
        subprocess.call("wscript DocxToPdf.vbs " + self.output_path)


    def replace_text_in_paragraph(self, paragraph, parameter):
        line = paragraph.runs
        for word in line:
            if word.text in parameter:
                value = parameter.get(word.text)
                if value is None:
                    value = ''
                word.text = word.text.replace(word.text, value)
            else:
                if PLACEHOLDER_OPENING in word.text and PLACEHOLDER_CLOSING in word.text:
                    self.delete_paragraph(paragraph)

    @staticmethod
    def delete_paragraph(paragraph):
        p = paragraph._element
        p.getparent().remove(p)
        paragraph._p = paragraph._element = None

    @staticmethod
    def build_file_name(parameter):
        return '{last_name}_{first_name}_{module}_{time}.docx'.format(
            last_name=parameter[PLACEHOLDER_OPENING + 'Nachname' + PLACEHOLDER_CLOSING],
            first_name=parameter[PLACEHOLDER_OPENING + 'Vorname' + PLACEHOLDER_CLOSING],
            module=parameter[PLACEHOLDER_OPENING + 'Modul' + PLACEHOLDER_CLOSING],
            time=time.time_ns())

# Output path is necessary as its later used to call the DocxToPDf script
pdf = PDFGenerator(
    template_file_name='C:/Users/Philipp/PycharmProjects/signer/Einzelzertifikat_final_mit Serienfeldern_abWS1819_mit_dig_unterschrift.docx',
    excel_file_name='Empfaenger.xlsx',
    output_path='C:/Users/Philipp/Downloads/Bachelor/docx-test/')
pdf.replace_parameters()
pdf.convert_docx_to_pdf()

